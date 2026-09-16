#!/usr/bin/env python3
"""
Terminal certificate generator for ITF National Age Group Triathlon.

Subcommands:
  merit          Generate merit certificates from the Podium sheet
  participation  Generate participation certificates from Overall results
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Tuple

import pandas as pd
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

# ---------------------------------------------------------------------------
# Paths / defaults
# ---------------------------------------------------------------------------

APP_DIR = Path(__file__).resolve().parent
FONTS_DIR = APP_DIR / "fonts"
DEFAULT_EXCEL = APP_DIR / "National Pune Results.xlsx"
DEFAULT_MERIT_TEMPLATE = APP_DIR / "new_merit_template.png"
DEFAULT_PART_TEMPLATE = APP_DIR / "new_participent_template.png"
DEFAULT_MERIT_SHEET = "Podium "
DEFAULT_PART_SHEET = "Overal results"
DEFAULT_MERIT_OUTPUT = APP_DIR / "output" / "merit"
DEFAULT_PART_OUTPUT = APP_DIR / "output" / "participation"
SEQUENCE_FILE = APP_DIR / "output" / ".next_cert_number"

CERTIFICATE_PREFIX = "ITF-"
CERTIFICATE_PAD = 2

# Body field text — navy to match templates
CERTIFICATE_TEXT_COLOR = (26, 43, 76)
# Certificate serial — dark navy for strong visibility
CERTIFICATE_NO_COLOR = (26, 43, 76)

FONT_CANDIDATES = {
    # Match the template's bold sans-serif look for body fields.
    "bold": [
        FONTS_DIR / "Arial Bold.ttf",
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
        "/Library/Fonts/Arial Bold.ttf",
        FONTS_DIR / "Times New Roman Bold.ttf",
        "/System/Library/Fonts/Supplemental/Times New Roman Bold.ttf",
    ],
    "regular": [
        FONTS_DIR / "Arial Bold.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        FONTS_DIR / "Times New Roman.ttf",
        "/System/Library/Fonts/Supplemental/Times New Roman.ttf",
    ],
}

# Field positions for 6400x4520 templates.
# y is the dashed underline target; text sits just above it.
_SHARED_6400_FIELD_CONFIG: Dict[str, Dict[str, Any]] = {
    "certificate_no": {
        "label": "CERTIFICATE_NO",
        "x": 2980,
        "y": 300,
        "max_width": 440,
        "font_size": 130,
        "min_font_size": 70,
        "align": "center",
        "font_weight": "bold",
        "color": CERTIFICATE_NO_COLOR,
        "calibration_color": (255, 0, 0),
    },
    "name": {
        "label": "NAME",
        "x": 1740,
        "y": 2692,
        "max_width": 3680,
        "font_size": 130,
        "min_font_size": 70,
        "align": "left",
        "font_weight": "bold",
        "calibration_color": (255, 0, 0),
    },
    "state": {
        "label": "STATE",
        "x": 1180,
        "y": 3016,
        "max_width": 1580,
        "font_size": 130,
        "min_font_size": 70,
        "align": "center",
        "font_weight": "bold",
        "calibration_color": (0, 128, 255),
    },
    "position": {
        "label": "POSITION",
        "x": 3270,
        "y": 3016,
        "max_width": 730,
        "font_size": 130,
        "min_font_size": 70,
        "align": "center",
        "font_weight": "bold",
        "calibration_color": (0, 180, 0),
    },
    "category": {
        "label": "CATEGORY",
        "x": 4720,
        "y": 3016,
        "max_width": 740,
        "font_size": 90,
        "min_font_size": 52,
        "align": "center",
        "font_weight": "bold",
        "max_lines": 2,
        "line_gap": 6,
        "calibration_color": (255, 128, 0),
    },
    "total_time": {
        "label": "TOTAL_TIME",
        "x": 3180,
        "y": 3340,
        "max_width": 1060,
        "font_size": 130,
        "min_font_size": 70,
        "align": "center",
        "font_weight": "bold",
        "calibration_color": (128, 0, 255),
    },
}

MERIT_FIELD_CONFIG: Dict[str, Dict[str, Any]] = {
    key: dict(value) for key, value in _SHARED_6400_FIELD_CONFIG.items()
}

PARTICIPATION_FIELD_CONFIG: Dict[str, Dict[str, Any]] = {
    key: dict(value) for key, value in _SHARED_6400_FIELD_CONFIG.items()
}

CALIBRATION_LABEL_OFFSET = (-4, -28)


# ---------------------------------------------------------------------------
# Exceptions
# ---------------------------------------------------------------------------


class CertificateGeneratorError(Exception):
    """Base exception for certificate generation errors."""


class MissingFileError(CertificateGeneratorError):
    """Raised when a required file is missing."""


class MissingColumnsError(CertificateGeneratorError):
    """Raised when required Excel columns are missing."""


class InvalidDataError(CertificateGeneratorError):
    """Raised when participant data is invalid."""


# ---------------------------------------------------------------------------
# Fonts / text helpers
# ---------------------------------------------------------------------------


def resolve_font_path(weight: str = "bold") -> Path:
    candidates = FONT_CANDIDATES.get(weight, FONT_CANDIDATES["bold"])
    for candidate in candidates:
        path = Path(candidate)
        if path.is_file():
            return path
    raise MissingFileError(
        f"Font file not found for weight '{weight}'. "
        f"Place Times New Roman fonts in {FONTS_DIR}/"
    )


def load_font(font_path: Path, size: int) -> ImageFont.FreeTypeFont:
    try:
        return ImageFont.truetype(str(font_path), size=size)
    except OSError as exc:
        raise CertificateGeneratorError(f"Unable to load font '{font_path}': {exc}") from exc


def fit_text_to_width(
    text: str,
    font_path: Path,
    max_width: int,
    initial_size: int,
    min_size: int = 12,
    *,
    max_lines: int = 1,
) -> Tuple[ImageFont.FreeTypeFont, List[str]]:
    """Fit text into max_width, optionally wrapping onto up to max_lines."""
    if not text:
        return load_font(font_path, initial_size), [""]

    size = initial_size
    while size >= min_size:
        font = load_font(font_path, size)
        if max_lines <= 1:
            width = font.getbbox(text)[2] - font.getbbox(text)[0]
            if width <= max_width:
                return font, [text]
        else:
            lines = wrap_text_to_width(text, font, max_width, max_lines)
            if lines is not None:
                return font, lines
        size -= 1

    font = load_font(font_path, min_size)
    if max_lines > 1:
        lines = wrap_text_to_width(text, font, max_width, max_lines)
        if lines is not None:
            return font, lines

    ellipsis = "..."
    truncated = text
    while truncated and (
        font.getbbox(truncated + ellipsis)[2] - font.getbbox(truncated + ellipsis)[0]
    ) > max_width:
        truncated = truncated[:-1]
    fitted = (truncated + ellipsis) if truncated != text else truncated
    return font, [fitted]


def wrap_text_to_width(
    text: str,
    font: ImageFont.FreeTypeFont,
    max_width: int,
    max_lines: int,
) -> Optional[List[str]]:
    """Return wrapped lines if text fits within max_lines, else None."""
    words = text.split()
    if not words:
        return [""]

    lines: List[str] = []
    current = words[0]
    for word in words[1:]:
        candidate = f"{current} {word}"
        width = font.getbbox(candidate)[2] - font.getbbox(candidate)[0]
        if width <= max_width:
            current = candidate
        else:
            lines.append(current)
            current = word
            if len(lines) >= max_lines:
                return None
    lines.append(current)
    if len(lines) > max_lines:
        return None

    for line in lines:
        width = font.getbbox(line)[2] - font.getbbox(line)[0]
        if width > max_width:
            return None
    return lines


def sanitize_filename(name: str) -> str:
    cleaned = re.sub(r"[^\w\s-]", "", name.strip())
    cleaned = re.sub(r"\s+", "_", cleaned)
    cleaned = cleaned.strip("_")
    if not cleaned:
        raise InvalidDataError("Name is blank or invalid.")
    return cleaned


def format_certificate_number(seq: int) -> str:
    if seq < 1:
        raise InvalidDataError("Certificate sequence must be at least 1.")
    width = CERTIFICATE_PAD if seq < 100 else len(str(seq))
    return f"{CERTIFICATE_PREFIX}{seq:0{width}d}"


def format_position(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        raise InvalidDataError("Position is missing.")

    text = str(value).strip()
    if not text or text.lower() == "nan":
        raise InvalidDataError("Position is missing.")

    text = text.rstrip(".")
    if text.upper() in {"DSQ", "DNF", "DNS", "DQ", "DNP"}:
        raise InvalidDataError(f"Non-finishing result: {text}")

    try:
        number = int(float(text))
    except ValueError as exc:
        raise InvalidDataError(f"Invalid position: {value!r}") from exc

    if 10 <= (number % 100) <= 20:
        suffix = "th"
    else:
        suffix = {1: "st", 2: "nd", 3: "rd"}.get(number % 10, "th")
    return f"{number}{suffix}"


def format_total_time(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        raise InvalidDataError("Total Time is missing.")

    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # Excel stores some times as fraction of a day.
        if 0 < float(value) < 1:
            total_seconds = float(value) * 86400.0
            hours = int(total_seconds // 3600)
            minutes = int((total_seconds % 3600) // 60)
            seconds = total_seconds % 60
            if hours > 0:
                return f"{hours}:{minutes:02d}:{seconds:05.2f}"
            return f"{minutes}:{seconds:05.2f}"
        text = str(int(value)) if float(value).is_integer() else str(value)
    else:
        text = str(value).strip()

    if not text or text.lower() == "nan":
        raise InvalidDataError("Total Time is missing.")

    # Already formatted MM:SS.ms or H:MM:SS.ms
    if re.fullmatch(r"\d{1,2}:\d{2}(?:\.\d+)?", text) or re.fullmatch(
        r"\d{1,2}:\d{2}:\d{2}(?:\.\d+)?", text
    ):
        return text

    # Numeric string day-fraction
    try:
        frac = float(text)
        if 0 < frac < 1:
            return format_total_time(frac)
    except ValueError:
        pass

    return text


def clean_cell(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    text = str(value).strip()
    if text.lower() == "nan":
        return ""
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".")[0]
    return text


def format_display_name(value: Any) -> str:
    """Title-case a name: first letter of each word capital, rest lowercase."""
    text = clean_cell(value)
    if not text:
        return ""
    return " ".join(word.capitalize() for word in text.split())


def gender_label(value: Any) -> str:
    text = clean_cell(value).upper()
    if text in {"F", "FEMALE"}:
        return "Female"
    if text in {"M", "MALE"}:
        return "Male"
    if not text:
        return ""
    return text.title()


def is_category_banner(bib_value: Any, name_value: Any) -> bool:
    """Yellow banner rows have category text in Bib and no athlete Name."""
    name = clean_cell(name_value)
    if name:
        return False
    bib = clean_cell(bib_value)
    if not bib:
        return False
    try:
        float(bib.replace(",", ""))
        return False
    except ValueError:
        return True


# ---------------------------------------------------------------------------
# Excel loaders
# ---------------------------------------------------------------------------


def _read_sheet_records(excel_path: Path, sheet_name: str) -> List[Dict[str, Any]]:
    if not excel_path.is_file():
        raise MissingFileError(f"Excel file not found: {excel_path}")

    try:
        workbook = load_workbook(excel_path, data_only=True, read_only=True)
    except Exception as exc:
        raise CertificateGeneratorError(f"Failed to open Excel file: {exc}") from exc

    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise MissingFileError(
            f"Sheet {sheet_name!r} not found. Available: {', '.join(workbook.sheetnames)}"
        )

    sheet = workbook[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()

    if not rows:
        raise InvalidDataError(f"Sheet {sheet_name!r} is empty.")

    headers = [clean_cell(h) for h in rows[0]]
    records: List[Dict[str, Any]] = []
    for row_number, row in enumerate(rows[1:], start=2):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        record: Dict[str, Any] = {"_row": row_number}
        for idx, header in enumerate(headers):
            if not header:
                continue
            record[header] = row[idx] if idx < len(row) else None
        records.append(record)
    return records


def load_merit_records(excel_path: Path, sheet_name: str = DEFAULT_MERIT_SHEET) -> List[Dict[str, str]]:
    """Parse Podium sheet: category banners + athlete rows."""
    records = _read_sheet_records(excel_path, sheet_name)

    required = ["Bib", "Name", "State", "Total Time", "Position"]
    # Headers may exist even if first data row is a banner
    sample_keys = set()
    for record in records[:5]:
        sample_keys.update(k for k in record if not k.startswith("_"))
    missing = [col for col in required if col not in sample_keys]
    if missing:
        raise MissingColumnsError(f"Missing required columns: {', '.join(missing)}")

    athletes: List[Dict[str, str]] = []
    current_category = ""

    for record in records:
        bib = record.get("Bib")
        name = record.get("Name")

        if is_category_banner(bib, name):
            current_category = clean_cell(bib)
            continue

        athlete_name = format_display_name(name)
        if not athlete_name:
            continue

        if not current_category:
            # Fallback: Category + Gender columns when banner missing
            cat = clean_cell(record.get("Category"))
            gen = gender_label(record.get("Gender"))
            current_category = f"{cat} - {gen}".strip(" -") if cat else ""

        try:
            athletes.append(
                {
                    "name": athlete_name,
                    "state": clean_cell(record.get("State")),
                    "position": format_position(record.get("Position")),
                    "category": current_category,
                    "total_time": format_total_time(record.get("Total Time")),
                    "source_row": str(record.get("_row", "")),
                }
            )
        except InvalidDataError as exc:
            print(
                f"Skipped podium row {record.get('_row')}: {exc}",
                file=sys.stderr,
            )

    if not athletes:
        raise InvalidDataError("No valid merit/podium athlete rows found.")
    return athletes


def load_participation_records(
    excel_path: Path,
    sheet_name: str = DEFAULT_PART_SHEET,
) -> List[Dict[str, str]]:
    """Parse Overall results sheet into certificate records."""
    records = _read_sheet_records(excel_path, sheet_name)

    required = ["Name", "State", "Category", "Gender", "Total Time", "Position"]
    sample_keys = set()
    for record in records[:5]:
        sample_keys.update(k for k in record if not k.startswith("_"))
    missing = [col for col in required if col not in sample_keys]
    if missing:
        raise MissingColumnsError(f"Missing required columns: {', '.join(missing)}")

    athletes: List[Dict[str, str]] = []
    for record in records:
        athlete_name = clean_cell(record.get("Name"))
        if not athlete_name:
            continue

        total_raw = record.get("Total Time")
        if total_raw is None or (isinstance(total_raw, float) and pd.isna(total_raw)) or clean_cell(total_raw) == "":
            print(
                f"Skipped overall row {record.get('_row')} ({athlete_name}): Total Time missing",
                file=sys.stderr,
            )
            continue

        category = clean_cell(record.get("Category"))
        gender = gender_label(record.get("Gender"))
        category_label = f"{category} - {gender}".strip(" -") if gender else category

        try:
            athletes.append(
                {
                    "name": athlete_name,
                    "state": clean_cell(record.get("State")),
                    "position": format_position(record.get("Position")),
                    "category": category_label,
                    "total_time": format_total_time(total_raw),
                    "source_row": str(record.get("_row", "")),
                }
            )
        except InvalidDataError as exc:
            print(
                f"Skipped overall row {record.get('_row')}: {exc}",
                file=sys.stderr,
            )

    if not athletes:
        raise InvalidDataError("No valid participation rows found.")
    return athletes


# ---------------------------------------------------------------------------
# Sequence file
# ---------------------------------------------------------------------------


def read_next_sequence(default: int = 1) -> int:
    if not SEQUENCE_FILE.is_file():
        return default
    try:
        value = int(SEQUENCE_FILE.read_text(encoding="utf-8").strip())
        return value if value >= 1 else default
    except ValueError:
        return default


def write_next_sequence(next_seq: int) -> None:
    SEQUENCE_FILE.parent.mkdir(parents=True, exist_ok=True)
    SEQUENCE_FILE.write_text(str(next_seq), encoding="utf-8")


# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------


def draw_field(
    draw: ImageDraw.ImageDraw,
    field_key: str,
    text: str,
    field_config: Dict[str, Dict[str, Any]],
    calibration: bool = False,
) -> None:
    config = field_config[field_key]
    x = config["x"]
    # y is the dashed underline / baseline target so font-size changes stay aligned
    underline_y = config["y"]
    max_width = config["max_width"]
    font_size = config["font_size"]
    min_font_size = config.get("min_font_size", 12)
    align = config.get("align", "left")
    text_color = config.get("color", CERTIFICATE_TEXT_COLOR)
    font_weight = config.get("font_weight", "bold")
    max_lines = int(config.get("max_lines", 1))
    line_gap = int(config.get("line_gap", 8))
    field_font_path = resolve_font_path(font_weight)

    cover = config.get("cover")
    if cover:
        draw.rectangle(
            [
                cover["x"],
                cover["y"],
                cover["x"] + cover["width"],
                cover["y"] + cover["height"],
            ],
            fill=tuple(cover.get("fill", (255, 255, 255))),
        )

    font, lines = fit_text_to_width(
        text=text,
        font_path=field_font_path,
        max_width=max_width,
        initial_size=font_size,
        min_size=min_font_size,
        max_lines=max_lines,
    )

    line_bboxes = [font.getbbox(line) for line in lines]
    line_heights = [bbox[3] - bbox[1] for bbox in line_bboxes]
    block_height = sum(line_heights) + line_gap * (len(lines) - 1)

    # Keep the bottom of the last line just above the dashed underline.
    block_bottom = underline_y - 6
    cursor_y = block_bottom - block_height

    ink_top = cursor_y
    ink_bottom = block_bottom

    for line, bbox, height in zip(lines, line_bboxes, line_heights):
        text_width = bbox[2] - bbox[0]
        if align == "center":
            draw_x = x + (max_width - text_width) // 2
        elif align == "right":
            draw_x = x + max_width - text_width
        else:
            draw_x = x

        # Align this line's ink top to cursor_y
        draw_y = cursor_y - bbox[1]
        draw.text((draw_x, draw_y), line, font=font, fill=text_color)
        cursor_y += height + line_gap

    if calibration:
        box_color = config.get("calibration_color", (255, 0, 0))
        draw.rectangle(
            [x, ink_top - 6, x + max_width, ink_bottom + 6],
            outline=box_color,
            width=3,
        )
        label = config.get("label", field_key.upper())
        label_font = load_font(resolve_font_path("regular"), 20)
        draw.text(
            (x + CALIBRATION_LABEL_OFFSET[0], ink_top + CALIBRATION_LABEL_OFFSET[1]),
            f"{label} ({x},{underline_y})",
            font=label_font,
            fill=box_color,
        )


def draw_coordinate_guides(
    draw: ImageDraw.ImageDraw,
    image_size: Tuple[int, int],
    font_path: Path,
) -> None:
    width, height = image_size
    guide_color = (200, 200, 200)
    label_font = load_font(font_path, 16)
    for x in range(0, width, 250):
        draw.line([(x, 0), (x, height)], fill=guide_color, width=1)
        draw.text((x + 4, 4), str(x), font=label_font, fill=guide_color)
    for y in range(0, height, 250):
        draw.line([(0, y), (width, y)], fill=guide_color, width=1)
        draw.text((4, y + 4), str(y), font=label_font, fill=guide_color)


def generate_certificate(
    template: Image.Image,
    field_values: Dict[str, str],
    field_config: Dict[str, Dict[str, Any]],
    calibration: bool = False,
) -> Image.Image:
    certificate = template.copy()
    draw = ImageDraw.Draw(certificate)

    if calibration:
        draw_coordinate_guides(draw, certificate.size, resolve_font_path("regular"))

    for field_key in field_config:
        draw_field(
            draw=draw,
            field_key=field_key,
            text=field_values.get(field_key, ""),
            field_config=field_config,
            calibration=calibration,
        )
    return certificate


def record_to_field_values(record: Dict[str, str], certificate_number: str) -> Dict[str, str]:
    if not record.get("name"):
        raise InvalidDataError("Name is missing.")
    if not record.get("state"):
        raise InvalidDataError("State is missing.")
    if not record.get("category"):
        raise InvalidDataError("Category is missing.")
    return {
        "certificate_no": certificate_number,
        "name": record["name"],
        "state": record["state"],
        "position": record["position"],
        "category": record["category"],
        "total_time": record["total_time"],
    }


# ---------------------------------------------------------------------------
# Batch generation
# ---------------------------------------------------------------------------


def generate_certificates(
    records: List[Dict[str, str]],
    template_path: Path,
    output_dir: Path,
    field_config: Dict[str, Dict[str, Any]],
    *,
    start_seq: int = 1,
    calibrate: bool = False,
) -> Tuple[int, int, int]:
    """
    Generate certificates for all records.

    Returns (generated_count, skipped_count, next_sequence).
    """
    if not template_path.is_file():
        raise MissingFileError(f"Template image not found: {template_path}")

    output_dir.mkdir(parents=True, exist_ok=True)
    resolve_font_path("bold")

    try:
        template = Image.open(template_path)
    except Exception as exc:
        raise CertificateGeneratorError(f"Failed to open template image: {exc}") from exc

    if calibrate:
        record = records[0]
        cert_no = format_certificate_number(start_seq)
        field_values = record_to_field_values(record, cert_no)
        preview = generate_certificate(
            template=template,
            field_values=field_values,
            field_config=field_config,
            calibration=True,
        )
        preview_path = output_dir / "calibration_preview.png"
        preview.save(preview_path, format="PNG")
        print(f"Generated:\n{preview_path.name}")
        print("\nCalibration preview created.")
        print("Adjust FIELD_CONFIG coordinates if needed, then re-run without --calibrate.")
        return 1, 0, start_seq

    generated = 0
    skipped = 0
    seq = start_seq
    total = len(records)

    for index, record in enumerate(records, start=1):
        try:
            cert_no = format_certificate_number(seq)
            field_values = record_to_field_values(record, cert_no)
            filename = f"{cert_no}_{sanitize_filename(field_values['name'])}.png"
            output_path = output_dir / filename

            certificate = generate_certificate(
                template=template,
                field_values=field_values,
                field_config=field_config,
                calibration=False,
            )
            certificate.save(output_path, format="PNG")
            generated += 1
            print(f"Generated ({index}/{total}):\n{filename}")
            seq += 1
        except InvalidDataError as exc:
            skipped += 1
            print(
                f"Skipped row {record.get('source_row', index)}: {exc}",
                file=sys.stderr,
            )

    return generated, skipped, seq


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Generate ITF merit and participation certificates from Excel.",
    )
    sub = parser.add_subparsers(dest="command", required=True)

    def add_common(cmd: argparse.ArgumentParser) -> None:
        cmd.add_argument("--excel", type=Path, default=DEFAULT_EXCEL, help="Path to Excel workbook")
        cmd.add_argument("--sheet", type=str, help="Worksheet name override")
        cmd.add_argument("--template", type=Path, help="Certificate template PNG override")
        cmd.add_argument("--output", type=Path, help="Output directory override")
        cmd.add_argument(
            "--start",
            type=int,
            default=None,
            help="Starting certificate sequence number (default: 1 for merit, "
            "or value from output/.next_cert_number for participation)",
        )
        cmd.add_argument(
            "--calibrate",
            action="store_true",
            help="Generate one calibration preview with field boxes",
        )

    merit = sub.add_parser("merit", help="Generate merit certificates from Podium sheet")
    add_common(merit)

    participation = sub.add_parser(
        "participation",
        help="Generate participation certificates from Overall results sheet",
    )
    add_common(participation)

    return parser


def run_command(args: argparse.Namespace) -> int:
    if args.command == "merit":
        excel_path = args.excel
        sheet = args.sheet or DEFAULT_MERIT_SHEET
        template_path = args.template or DEFAULT_MERIT_TEMPLATE
        output_dir = args.output or DEFAULT_MERIT_OUTPUT
        field_config = MERIT_FIELD_CONFIG
        start_seq = args.start if args.start is not None else 1
        loader: Callable[[Path, str], List[Dict[str, str]]] = load_merit_records
        label = "Merit"
    elif args.command == "participation":
        excel_path = args.excel
        sheet = args.sheet or DEFAULT_PART_SHEET
        template_path = args.template or DEFAULT_PART_TEMPLATE
        output_dir = args.output or DEFAULT_PART_OUTPUT
        field_config = PARTICIPATION_FIELD_CONFIG
        start_seq = args.start if args.start is not None else read_next_sequence(1)
        loader = load_participation_records
        label = "Participation"
    else:
        raise CertificateGeneratorError(f"Unknown command: {args.command}")

    if args.start is not None and args.start < 1:
        raise CertificateGeneratorError("--start must be at least 1.")

    print(f"{label} certificate generation")
    print(f"Excel: {excel_path}")
    print(f"Sheet: {sheet}")
    print(f"Template: {template_path}")
    print(f"Output: {output_dir}")
    print(f"Starting certificate number: {format_certificate_number(start_seq)}")
    print("---")

    records = loader(excel_path, sheet)
    print(f"Records loaded: {len(records)}")

    generated, skipped, next_seq = generate_certificates(
        records=records,
        template_path=template_path,
        output_dir=output_dir,
        field_config=field_config,
        start_seq=start_seq,
        calibrate=args.calibrate,
    )

    if not args.calibrate:
        write_next_sequence(next_seq)
        print("\n---")
        print(f"Certificates Generated: {generated}")
        if skipped:
            print(f"Skipped: {skipped}")
        print(f"Output Folder: {output_dir}")
        print(f"Next certificate number: {format_certificate_number(next_seq)}")
        print("---------------------")

    return 0 if generated > 0 else 1


def main(argv: Optional[List[str]] = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        return run_command(args)
    except CertificateGeneratorError as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1
    except Exception as exc:
        print(f"Unexpected error: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
